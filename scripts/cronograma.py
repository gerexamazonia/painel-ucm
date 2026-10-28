# Lê o "Cronograma UcM (Demanda Presidência).xlsx" e extrai, por contrato, as datas de
# Início e Fim previstos de cada atividade (colunas G e H das abas 0072026 ... 0222026).
# Não lê nomes de pessoas (gestores) nem valores financeiros.
#
# Uso: python3 scripts/cronograma.py <cronograma.xlsx> <saida.json> "<data de modificação do arquivo>"
import datetime as dt
import json
import re
import sys

import openpyxl

src, out, mod = sys.argv[1], sys.argv[2], (sys.argv[3] if len(sys.argv) > 3 else '')
wb = openpyxl.load_workbook(src, data_only=True)
wbf = openpyxl.load_workbook(src, data_only=False)  # para saber se a célula ainda é fórmula (data do edital)
EPOCH = dt.datetime(1899, 12, 30)


def data(v):
    """Aceita datetime, número de série do Excel ou texto dd/mm/aaaa; devolve 'AAAA-MM-DD' ou None."""
    if v is None or v == '':
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (EPOCH + dt.timedelta(days=float(v))).date().isoformat()
    m = re.match(r'^\s*(\d{1,2})/(\d{1,2})/(\d{2,4})\s*$', str(v))
    if m:
        d, mo, y = map(int, m.groups())
        y += 2000 if y < 100 else 0
        try:
            return dt.date(y, mo, d).isoformat()
        except ValueError:
            return None
    return None


contratos, avisos = [], []
for nome in wb.sheetnames:
    if not re.fullmatch(r'\d{7}', nome):
        continue
    ws = wb[nome]
    wf = wbf[nome]
    base = ws['E6'].value
    contrato = str(ws['C2'].value or '').strip() or f'{nome[:3]}/{nome[3:]}'
    # localiza o cabeçalho (linha com "Início" na coluna G) para tolerar linhas inseridas
    head = next((r for r in range(1, 20) if 'In' in str(ws.cell(r, 7).value or '') and 'Fim' in str(ws.cell(r, 8).value or '')), 8)
    ativ = []
    for r in range(head + 1, head + 40):
        n = ws.cell(r, 1).value
        if not isinstance(n, (int, float)):
            if 'TOTAL' in str(ws.cell(r, 2).value or '').upper():
                break
            continue
        def val(col):
            v, f = ws.cell(r, col).value, wf.cell(r, col).value
            if isinstance(f, str) and f.startswith('='):
                m = re.search(r'EDATE\(\$E\$6,\s*(\d+)\)', f)
                if v is None and m and isinstance(base, dt.datetime):  # arquivo salvo sem valores calculados
                    mo = base.month - 1 + int(m.group(1)); y = base.year + mo // 12; mo = mo % 12 + 1
                    import calendar
                    v = base.replace(year=y, month=mo, day=min(base.day, calendar.monthrange(y, mo)[1]))
                return v, 'edital'
            return v, ('plano' if v not in (None, '') else '')
        (vi, fi), (vf, ff) = val(7), val(8)
        ini, fim = data(vi), data(vf)
        if vi not in (None, '') and ini is None:
            avisos.append(f'{contrato} atividade {int(n)}: Início previsto ilegível ({vi!r})')
        if vf not in (None, '') and fim is None:
            avisos.append(f'{contrato} atividade {int(n)}: Fim previsto ilegível ({vf!r})')
        if ini and fim and fim < ini:
            avisos.append(f'{contrato} atividade {int(n)}: Fim previsto ({fim}) antes do Início ({ini})')
        ativ.append({'n': int(n), 'rota': str(ws.cell(r, 2).value or '').split(' - ')[0].strip(),
                     'atividade': str(ws.cell(r, 3).value or '').strip(), 'inicio': ini, 'fim': fim,
                     'fonte': 'plano' if 'plano' in (fi, ff) else 'edital',
                     'status': str(ws.cell(r, 10).value or '').strip()})
    if not any(a['fim'] for a in ativ if a['n'] <= 4):
        avisos.append(f'{contrato}: nenhuma data de Fim previsto nas atividades da Rota 1 (o painel usará OS + 5 meses)')
    # linha 7: Plano de Ação (SEI) e data da atualização
    sei = str(ws['E7'].value or '').strip() if str(ws['A7'].value or '').startswith('Plano') else ''
    atualizado = data(ws['L7'].value) if str(ws['I7'].value or '').startswith('Datas') else None
    di = data(ws['E6'].value)
    for a in ativ:
        contratos.append(dict({'chave': f"{contrato}|{a['n']}", 'contrato': contrato, 'data_inicio': di, 'plano_sei': sei, 'datas_atualizadas_em': atualizado}, **a))

json.dump({'arquivo_modificado_em': mod, 'linhas': contratos, 'avisos': avisos}, open(out, 'w'), ensure_ascii=False)
print('ok cronograma', len({c['contrato'] for c in contratos}), 'contratos', len(contratos), 'linhas', len(avisos), 'avisos')
for a in avisos:
    print('AVISO', a)
